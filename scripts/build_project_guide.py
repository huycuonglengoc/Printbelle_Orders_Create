from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Printbelle_Project_Guide.docx"

BLACK = "000000"
TEXT = "252A27"
MUTED = "66706A"
GREEN = "285F50"
LIGHT_GREEN = "EAF2EE"
LIGHT_GRAY = "F3F5F2"
BORDER = "D9DED9"
ORANGE = "D85D3C"
NUMBER_COUNTER = 0


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=BORDER, size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def set_run_font(run, name="Aptos", size=10.5, color=TEXT, bold=False, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def style_paragraph(paragraph, before=0, after=6, line=1.15):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_text(paragraph, text, **kwargs):
    run = paragraph.add_run(text)
    set_run_font(run, **kwargs)
    return run


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    style_paragraph(p, before=14 if level == 1 else 8, after=6, line=1.05)
    add_text(p, text, size=16 if level == 1 else 12.5, color=BLACK, bold=True)
    return p


def add_body(doc, text, after=7):
    p = doc.add_paragraph(style="Body Text")
    style_paragraph(p, after=after)
    add_text(p, text, size=10.5, color=TEXT)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    style_paragraph(p, after=3)
    add_text(p, text, size=10.3, color=TEXT)
    return p


def add_number(doc, text):
    global NUMBER_COUNTER
    NUMBER_COUNTER += 1
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    style_paragraph(p, after=3)
    add_text(p, f"{NUMBER_COUNTER}. ", size=10.3, color=TEXT)
    add_text(p, text, size=10.3, color=TEXT)
    return p


def reset_numbering():
    global NUMBER_COUNTER
    NUMBER_COUNTER = 0


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.right_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.0
    set_cell_shading(p._p, LIGHT_GRAY) if False else None
    for index, line in enumerate(text.strip("\n").splitlines()):
        run = p.add_run(line)
        set_run_font(run, name="Courier New", size=8.2, color=TEXT)
        if index < len(text.strip("\n").splitlines()) - 1:
            run.add_break()
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        if widths:
            cell.width = Inches(widths[i])
        set_cell_shading(cell, GREEN)
        set_cell_borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        style_paragraph(p, after=0, line=1.0)
        add_text(p, header, size=9, color="FFFFFF", bold=True)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            if widths:
                cell.width = Inches(widths[i])
            set_cell_shading(cell, "FFFFFF" if row_index % 2 == 0 else LIGHT_GRAY)
            set_cell_borders(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            style_paragraph(p, after=0, line=1.1)
            add_text(p, str(value), size=9, color=TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_link_line(doc, label, url):
    p = doc.add_paragraph(style="Body Text")
    style_paragraph(p, after=4)
    add_text(p, f"{label}: ", size=10.3, color=TEXT, bold=True)
    add_text(p, url, size=10.3, color=GREEN)
    return p


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = styles[name]
        style.font.name = "Aptos Display" if name == "Title" else "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.color.rgb = RGBColor.from_string(BLACK)
        if name == "Title":
            p_pr = style._element.get_or_add_pPr()
            borders = p_pr.find(qn("w:pBdr"))
            if borders is not None:
                p_pr.remove(borders)
    if "Code" not in [s.name for s in styles]:
        code_style = styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
        code_style.font.name = "Courier New"
        code_style._element.rPr.rFonts.set(qn("w:ascii"), "Courier New")
        code_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Courier New")
        code_style.font.size = Pt(8.7)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(paragraph, "Printbelle Order Desk  ·  ", size=8, color=MUTED)
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def build():
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    add_page_number(section.footer.paragraphs[0])

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style_paragraph(title, before=24, after=8, line=1.0)
    remove_paragraph_borders(title)
    add_text(title, "Printbelle Order Desk Project Guide", size=27, color=BLACK, bold=True)
    subtitle = doc.add_paragraph()
    style_paragraph(subtitle, after=14, line=1.1)
    add_text(subtitle, "Tài liệu tóm tắt dự án và hướng dẫn chạy local và triển khai Railway", size=12, color=MUTED)
    meta = doc.add_paragraph()
    style_paragraph(meta, after=18, line=1.0)
    add_text(meta, "Phiên bản tài liệu 1.0  ·  Ngày 19 tháng 10 năm 2026", size=9.5, color=GREEN, bold=True)

    add_body(doc, "Printbelle Order Desk là web app Node.js/Express giúp đọc packing slip PDF, ghép artwork theo khách hàng, ánh xạ SPU và biến thể, kiểm tra payload rồi chuẩn bị gửi đơn sang Printbelle. Tài liệu này tập trung vào cách hiểu hệ thống và cách đưa app chạy được từ máy local lên Railway.", after=12)
    add_table(doc, ["Mục tiêu", "Kết quả"], [
        ("Giảm thao tác đặt đơn", "Tự đọc PDF, ghép artwork và dựng payload."),
        ("Giảm lỗi mapping", "Lưu mapping SPU, kích cỡ, màu và variant trong catalog config."),
        ("Không gửi nhầm đơn", "Preview và dry-run trước khi bật live API."),
        ("Dễ triển khai", "Có Dockerfile, PORT binding và Railway healthcheck."),
    ], widths=[1.75, 4.9])

    add_heading(doc, "1 Tổng quan chức năng")
    add_body(doc, "Luồng xử lý bắt đầu từ một PDF packing slip và các file artwork. App đọc từng trang PDF thành một order, nhận diện thông tin giao hàng và các item, sau đó ghép artwork theo tên khách hàng trong filename. Người dùng có thể sửa SPU, kích cỡ, màu và địa chỉ trước khi tạo preview.")
    add_table(doc, ["Chức năng", "Mô tả"], [
        ("Đọc PDF", "Đọc nhiều trang trong input.pdf; mỗi trang được xử lý thành một order."),
        ("Đọc item", "Nhận diện số lượng, SKU nguồn, product code, màu, kích cỡ và personalization."),
        ("Ghép artwork", "Ghép file theo tên khách hàng và sắp xếp theo số thứ tự tăng dần trong tên file."),
        ("SPU dropdown", "Tự nhận diện bằng alias/SKU prefix hoặc cho phép chọn/nhập SPU trong giao diện."),
        ("Mapping config", "Có trang /settings riêng để chỉnh SPU, prefix và kích cỡ; lưu dùng chung trên server."),
        ("Review và skip", "Order lỗi được đánh dấu và bỏ qua khi submit; giao diện hiển thị lý do."),
        ("Submit API", "Kiểm tra orderId trùng rồi gửi từng order hợp lệ qua adapter Printbelle."),
        ("Dry-run", "Dựng payload mà chưa gọi endpoint tạo đơn thật; mặc định được bật để an toàn."),
    ], widths=[1.55, 5.1])

    add_heading(doc, "2 Quy tắc xử lý nhiều item và nhiều artwork")
    add_body(doc, "App phân biệt hai trường hợp để tránh tạo dư line item hoặc ghép sai artwork.")
    add_heading(doc, "2.1 Một item nhiều artwork", level=2)
    add_bullet(doc, "Giữ lại một line item duy nhất.")
    add_bullet(doc, "Nếu có nhiều artwork, quantity được đặt theo số artwork. Ví dụ 37 artwork sẽ thành quantity 37.")
    add_bullet(doc, "Tất cả artwork vẫn hiển thị trên giao diện để đối chiếu. Theo schema adapter hiện tại, URL artwork đầu tiên được đưa vào field artwork của API item.")
    add_heading(doc, "2.2 Nhiều item mỗi item một artwork", level=2)
    add_bullet(doc, "Artwork được sắp xếp theo số thứ tự đầu filename, ví dụ 1_, 2_, 3_.")
    add_bullet(doc, "Artwork thứ nhất ghép với item thứ nhất, artwork thứ hai ghép với item thứ hai.")
    add_bullet(doc, "Số artwork phải đúng bằng số item và mỗi item phải có quantity bằng 1.")
    add_bullet(doc, "Nếu thiếu, thừa hoặc không đủ điều kiện ghép, order được đánh dấu lỗi và không được gửi.")

    add_heading(doc, "3 Cấu trúc thư mục")
    add_table(doc, ["Đường dẫn", "Vai trò"], [
        ("src/server.js", "Express server, upload, API routes, healthcheck và shutdown."),
        ("src/services/pdfParser.js", "Đọc text từ PDF bằng pdfjs-dist."),
        ("src/services/orderParser.js", "Parse thông tin shipping và item từ các dòng PDF."),
        ("src/services/artworkAssignment.js", "Áp dụng quy tắc bulk item hoặc one-artwork-per-item."),
        ("src/services/orderBuilder.js", "Ánh xạ config và dựng payload Printbelle."),
        ("src/services/printbelleClient.js", "Adapter gọi API Printbelle."),
        ("data/catalog.json", "Cấu hình mặc định/fallback của SPU, sizeMappings, colorMappings và variantMappings."),
        ("public/", "Giao diện đặt hàng và trang /settings bằng HTML, CSS, JavaScript thuần."),
        ("Dockerfile / railway.json", "Cấu hình build, healthcheck và restart trên Railway."),
    ], widths=[2.35, 4.3])

    add_heading(doc, "4 Chạy ở local")
    add_body(doc, "Yêu cầu: Node.js 20 trở lên và npm. Không cần cài database cho bản hiện tại.")
    reset_numbering()
    add_number(doc, "Mở Terminal tại thư mục project.")
    add_code(doc, "cd /Users/huycuong/Desktop/Printbelle")
    add_number(doc, "Cài dependency.")
    add_code(doc, "npm install")
    add_number(doc, "Tạo file biến môi trường local.")
    add_code(doc, "cp .env.example .env")
    add_number(doc, "Chạy test parser.")
    add_code(doc, "npm test")
    add_number(doc, "Khởi động app.")
    add_code(doc, "npm start")
    add_number(doc, "Mở trình duyệt tại địa chỉ sau.")
    add_code(doc, "http://localhost:3000")
    add_body(doc, "Trong giao diện, chọn input.pdf và các artwork PNG/JPG/WebP. Nếu cần sửa SPU/kích cỡ, mở trang /settings trước hoặc trong tab khác, bấm Lưu cấu hình rồi quay lại trang đặt hàng và chọn Tạo preview payload. Khi PRINTBELLE_DRY_RUN=true, bước cuối chỉ mô phỏng và không gửi order thật.")

    add_heading(doc, "5 Cấu hình môi trường")
    add_table(doc, ["Biến", "Bắt buộc", "Ý nghĩa"], [
        ("PORT", "Railway tự cấp", "Port server lắng nghe; app bind trên 0.0.0.0."),
        ("PRINTBELLE_API_BASE_URL", "Không", "Mặc định https://fulfillment.printbelle.com."),
        ("PRINTBELLE_API_TOKEN", "Khi gọi API", "Token được gửi qua header PB-ACCESS-TOKEN."),
        ("PRINTBELLE_CREATE_ORDER_PATH", "Khi live", "Endpoint tạo order; mặc định /prod_api/api/orders/v2/create và có thể override."),
        ("PRINTBELLE_LIST_ORDERS_PATH", "Khi live", "Endpoint kiểm tra orderId trùng; mặc định /prod_api/api/orders/v2/all."),
        ("PRINTBELLE_API_TIMEOUT_MS", "Không", "Timeout cho mỗi request Printbelle; mặc định 30000 ms."),
        ("PRINTBELLE_DRY_RUN", "Không", "Mặc định true; đặt false chỉ sau khi kiểm tra endpoint và payload."),
        ("ARTWORK_BASE_URL", "Khi live", "Domain public để Printbelle fetch artwork."),
        ("UPLOAD_DIR", "Không", "Thư mục upload tạm; mặc định ./var/uploads."),
        ("CATALOG_CONFIG_PATH", "Không", "File cấu hình server; mặc định ./data/catalog.json. Railway nên đặt trên Volume."),
    ], widths=[2.15, 1.05, 3.45])

    add_heading(doc, "6 Mapping SPU và kích cỡ")
    add_body(doc, "Không cần chỉnh JSON thủ công. Mở trang /settings hoặc bấm Cấu hình SPU trên trang chính. Danh mục sản phẩm dùng để thêm tên hiển thị, mã SPU, SKU prefix và các alias trong PDF; phần Ánh xạ kích cỡ dùng để chuyển kích thước trong PDF sang giá trị menu/variant của Printbelle. Size mapping có thể dùng SKU prefix hoặc product code đầy đủ.")
    add_bullet(doc, "Bấm + Thêm SPU, nhập tên, mã SPU, SKU prefix và các từ khóa ngăn cách bằng dấu phẩy. Prefix GF sẽ khớp GF-10, GF-20… và cùng trỏ về một SPU.")
    add_bullet(doc, "Bấm + Thêm sản phẩm trong phần Ánh xạ kích cỡ, nhập SKU prefix hoặc mã sản phẩm trong PDF rồi thêm các cặp kích thước nguồn và kích thước Printbelle. Ví dụ prefix GF áp dụng cho GF-10, GF-20…; mã đầy đủ chỉ áp dụng cho SKU đó.")
    add_bullet(doc, "Bấm Lưu cấu hình để ghi lên server. Các trình duyệt khác mở cùng app sẽ dùng chung cấu hình.")
    add_bullet(doc, "data/catalog.json chỉ là cấu hình mặc định/fallback khi server chưa có file cấu hình riêng.")
    add_body(doc, "Nếu Printbelle yêu cầu productCode cụ thể, variantMappings vẫn là nhóm nâng cao trong file cấu hình. Key có thể dùng dạng PRODUCT|SIZE|COLOR. Với SPU và size, sau khi bấm Lưu cấu hình không cần restart server.")

    add_heading(doc, "7 Build và deploy trên Railway")
    add_body(doc, "Project đã có Dockerfile và railway.json. Railway sẽ build container Node 20, chạy npm start và kiểm tra endpoint /api/health.")
    add_heading(doc, "7.1 Deploy từ GitHub", level=2)
    reset_numbering()
    add_number(doc, "Push project lên GitHub, bao gồm package.json, package-lock.json, src, public, data, Dockerfile và railway.json.")
    add_number(doc, "Trong Railway, chọn New Project rồi Deploy from GitHub Repo.")
    add_number(doc, "Chọn repository Printbelle và chờ Railway detect Dockerfile.")
    add_number(doc, "Vào Variables và thêm các biến production ở bảng cấu hình.")
    add_number(doc, "Tạo public domain trong Settings Networking.")
    add_number(doc, "Đặt ARTWORK_BASE_URL bằng public domain vừa tạo, không thêm dấu slash cuối.")
    add_number(doc, "Kiểm tra log build và truy cập /api/health trước khi dùng giao diện.")
    add_heading(doc, "7.2 Các biến production tối thiểu", level=2)
    add_code(doc, '''NODE_ENV=production
PRINTBELLE_API_TOKEN=your_printbelle_token
PRINTBELLE_CREATE_ORDER_PATH=/prod_api/api/orders/v2/create
PRINTBELLE_LIST_ORDERS_PATH=/prod_api/api/orders/v2/all
PRINTBELLE_API_TIMEOUT_MS=30000
PRINTBELLE_DRY_RUN=true
ARTWORK_BASE_URL=https://your-app.up.railway.app
CATALOG_CONFIG_PATH=/data/catalog.json
UPLOAD_DIR=/data/uploads''')
    add_body(doc, "Chỉ đổi PRINTBELLE_DRY_RUN=false sau khi Printbelle xác nhận endpoint tạo đơn và schema request. Có thể dùng dry-run để kiểm tra payload trong production mà không phát sinh order.")
    add_heading(doc, "7.3 Lưu ý về file upload và config trên Railway", level=2)
    add_body(doc, "Cấu hình SPU/kích cỡ và artwork đều được lưu trên server. Railway filesystem có thể là ephemeral sau redeploy hoặc restart, vì vậy hãy tạo một Railway Volume mount tại /data nếu muốn giữ dữ liệu lâu dài. Đặt CATALOG_CONFIG_PATH=/data/catalog.json và UPLOAD_DIR=/data/uploads. Nếu không cần giữ artwork, chỉ cần đặt CATALOG_CONFIG_PATH trên Volume để bảo toàn mapping. Với artwork lâu dài, có thể thay storage local bằng S3 hoặc Cloudflare R2 và đặt ARTWORK_BASE_URL về URL ổn định.")

    add_heading(doc, "8 API và giới hạn hiện tại")
    add_body(doc, "Luồng submit thật được tách riêng trong printbelleClient.js và orderBuilder.js. Khi live, app dùng PB-ACCESS-TOKEN, kiểm tra orderId đã tồn tại qua endpoint danh sách đơn, bỏ qua duplicate, rồi gửi tuần tự từng order hợp lệ. Public document xác nhận token endpoint, header PB-ACCESS-TOKEN và endpoint tra cứu order; endpoint tạo order không hiển thị đầy đủ schema nên create path vẫn được cấu hình bằng biến môi trường.")
    add_link_line(doc, "Printbelle API document", "https://documenter.getpostman.com/view/25923095/2s93RL1vvB")
    add_bullet(doc, "Không commit file .env hoặc PRINTBELLE_API_TOKEN lên GitHub.")
    add_bullet(doc, "Luôn kiểm tra payload preview, SPU, size, color, địa chỉ và artwork trước khi bật live.")
    add_bullet(doc, "Order lỗi được bỏ qua theo từng order; order hợp lệ khác vẫn có thể tiếp tục submit.")
    add_bullet(doc, "Nếu account dùng create path hoặc field artwork khác, đổi biến môi trường hoặc cập nhật buildPrintbellePayload trong orderBuilder.js.")

    add_heading(doc, "9 Kiểm tra nhanh khi có lỗi")
    add_table(doc, ["Hiện tượng", "Cách kiểm tra"], [
        ("Không đọc được PDF", "Kiểm tra file là PDF text-based và chạy npm test để xem parser."),
        ("Không thấy artwork", "Tên khách hàng trong filename phải khớp tên trên PDF; kiểm tra cả dấu apostrophe và số thứ tự."),
        ("Order bị bỏ qua", "Xem phần lỗi màu đỏ dưới order và payload preview."),
        ("Railway không healthy", "Mở /api/health, kiểm tra PORT, log deploy và Dockerfile."),
        ("Printbelle không lấy được artwork", "Kiểm tra ARTWORK_BASE_URL là public HTTPS và file URL có thể truy cập."),
        ("API trả lỗi tạo order", "Để dry-run, kiểm tra create path và schema trong printbelleClient/orderBuilder."),
    ], widths=[2.1, 4.55])

    add_heading(doc, "10 Các lệnh thường dùng")
    add_code(doc, '''npm install        # cài dependency
npm test           # chạy test parser và artwork assignment
npm start          # chạy production server local
npm run dev        # chạy node watch trong lúc phát triển''')
    add_body(doc, "Tài liệu này mô tả trạng thái và cấu hình của project tại ngày 19 tháng 10 năm 2026. Khi Printbelle thay đổi schema, cập nhật đồng thời .env, README.md, orderBuilder.js và phần API adapter.", after=0)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
